import json
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


DATA_DIR = Path("endpoint-data")
OUTPUT_FILE = "Nanobricks_Endpoint_Assessment.xlsx"


# ============================================================
# COLORS
# ============================================================

HEADER_FILL = "1F4E78"
HEADER_FONT = "FFFFFF"

PASS_FILL = "C6EFCE"
REVIEW_FILL = "FFF2CC"
FAIL_FILL = "FFC7CE"
NEUTRAL_FILL = "E7E6E6"

THIN_BORDER = Border(
    bottom=Side(
        style="thin",
        color="D9E1F2"
    )
)


# ============================================================
# DATA HELPERS
# ============================================================

def load_reports():
    reports = []

    for file_path in sorted(DATA_DIR.glob("*-assessment.json")):
        try:
            with open(file_path, "r", encoding="utf-8") as file:
                report = json.load(file)

            reports.append(report)

        except (json.JSONDecodeError, OSError) as error:
            print(
                f"Skipping {file_path.name}: {error}"
            )

    return reports


def get_overall_status(summary):
    if summary.get("fail", 0) > 0:
        return "FAIL"

    if summary.get("review", 0) > 0:
        return "REVIEW"

    return "PASS"


def yes_no(value):
    if value is True:
        return "Yes"

    if value is False:
        return "No"

    return "Unknown"


def get_firewall_status(raw_data, profile_name):
    firewall = raw_data.get("firewall", {})

    if not firewall.get("available", False):
        return "Unknown"

    for profile in firewall.get("profiles", []):
        if profile.get("name") == profile_name:
            return (
                "Enabled"
                if profile.get("enabled") is True
                else "Disabled"
            )

    return "Unknown"


def get_enabled_users(raw_data):
    local_users = raw_data.get("local_users", {})

    users = local_users.get("users", [])

    enabled_users = [
        user.get("Name", "")
        for user in users
        if user.get("Enabled") is True
    ]

    return ", ".join(
        name for name in enabled_users if name
    )


def get_local_administrators(raw_data):
    local_users = raw_data.get("local_users", {})

    administrators = local_users.get(
        "administrators",
        []
    )

    admin_names = [
        admin.get("Name", "")
        for admin in administrators
        if admin.get("Name")
    ]

    return ", ".join(admin_names)


def format_defender_signature(value):
    """
    Convert PowerShell/.NET date values such as:
    /Date(1790656934000)/
    into:
    29-Sep-2026 04:42:14 UTC
    """

    if not value:
        return ""

    if not isinstance(value, str):
        return str(value)

    match = re.search(
        r"/Date\((-?\d+)(?:[+-]\d+)?\)/",
        value
    )

    if not match:
        return value

    try:
        milliseconds = int(match.group(1))

        signature_time = datetime.fromtimestamp(
            milliseconds / 1000,
            tz=timezone.utc
        )

        return signature_time.strftime(
            "%d-%b-%Y %H:%M:%S UTC"
        )

    except (ValueError, OSError, OverflowError):
        return value


# ============================================================
# EXCEL FORMATTING HELPERS
# ============================================================

def fill_cell(cell, color):
    cell.fill = PatternFill(
        fill_type="solid",
        fgColor=color
    )


def apply_status_color(cell):
    status = str(cell.value).strip().upper()

    if status == "PASS":
        fill_cell(cell, PASS_FILL)

    elif status == "REVIEW":
        fill_cell(cell, REVIEW_FILL)

    elif status == "FAIL":
        fill_cell(cell, FAIL_FILL)


def apply_yes_no_color(cell, no_is_review=False):
    value = str(cell.value).strip().lower()

    if value == "yes":
        fill_cell(cell, PASS_FILL)

    elif value == "no":
        fill_cell(
            cell,
            REVIEW_FILL if no_is_review else FAIL_FILL
        )

    elif value == "unknown":
        fill_cell(cell, REVIEW_FILL)


def apply_enabled_disabled_color(cell):
    value = str(cell.value).strip().lower()

    if value == "enabled":
        fill_cell(cell, PASS_FILL)

    elif value == "disabled":
        fill_cell(cell, FAIL_FILL)

    elif value == "unknown":
        fill_cell(cell, REVIEW_FILL)


def apply_ram_color(cell):
    try:
        ram = float(cell.value)

        if ram >= 8:
            fill_cell(cell, PASS_FILL)

        elif ram >= 4:
            fill_cell(cell, REVIEW_FILL)

        else:
            fill_cell(cell, FAIL_FILL)

    except (TypeError, ValueError):
        fill_cell(cell, REVIEW_FILL)


def apply_disk_color(cell):
    try:
        used_percent = float(cell.value)

        if used_percent <= 85:
            fill_cell(cell, PASS_FILL)

        elif used_percent <= 90:
            fill_cell(cell, REVIEW_FILL)

        else:
            fill_cell(cell, FAIL_FILL)

    except (TypeError, ValueError):
        fill_cell(cell, REVIEW_FILL)


def apply_updates_color(cell):
    try:
        pending = int(cell.value)

        if pending == 0:
            fill_cell(cell, PASS_FILL)

        elif pending <= 5:
            fill_cell(cell, REVIEW_FILL)

        else:
            fill_cell(cell, FAIL_FILL)

    except (TypeError, ValueError):
        fill_cell(cell, REVIEW_FILL)


def apply_os_color(cell):
    value = str(cell.value).strip().lower()

    if value == "windows 11":
        fill_cell(cell, PASS_FILL)

    elif value:
        fill_cell(cell, FAIL_FILL)

    else:
        fill_cell(cell, REVIEW_FILL)


def style_header(sheet):
    for cell in sheet[1]:
        fill_cell(cell, HEADER_FILL)

        cell.font = Font(
            color=HEADER_FONT,
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )


def autofit(sheet):
    for column_cells in sheet.columns:
        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            value = str(
                cell.value
                if cell.value is not None
                else ""
            )

            max_length = max(
                max_length,
                len(value)
            )

        sheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            42
        )


def format_sheet(sheet):
    style_header(sheet)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    sheet.row_dimensions[1].height = 32

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

            cell.border = THIN_BORDER

    autofit(sheet)


# ============================================================
# EXCEL GENERATOR
# ============================================================

def build_excel(reports):
    workbook = Workbook()

    inventory_sheet = workbook.active
    inventory_sheet.title = "Endpoint Inventory"

    summary_sheet = workbook.create_sheet(
        "Assessment Summary"
    )

    findings_sheet = workbook.create_sheet(
        "Findings"
    )

    users_sheet = workbook.create_sheet(
        "Local Users"
    )

    # ========================================================
    # ENDPOINT INVENTORY
    # ========================================================

    inventory_headers = [
        "Hostname",
        "Operating System",
        "OS Version",
        "Architecture",
        "Processor",
        "RAM (GB)",
        "Disk Total (GB)",
        "Disk Free (GB)",
        "Disk Used (%)",
        "Defender Antivirus",
        "Real-Time Protection",
        "Defender Signature",
        "Firewall Domain",
        "Firewall Private",
        "Firewall Public",
        "Pending Updates",
        "Office Installed",
        "Office Product",
        "Office Version",
        "Office Platform",
        "Enabled Local Users",
        "Local Administrators",
        "PASS",
        "REVIEW",
        "FAIL",
        "Total Checks",
        "Overall Status",
        "Collected At UTC"
    ]

    inventory_sheet.append(inventory_headers)

    for report in reports:
        endpoint = report.get("endpoint", {})
        raw_data = report.get("raw_data", {})
        summary = report.get("summary", {})
        agent = report.get("agent", {})

        defender = raw_data.get("defender", {})
        updates = raw_data.get("windows_update", {})
        office = raw_data.get("office", {})

        inventory_sheet.append(
            [
                endpoint.get("hostname", ""),
                endpoint.get("operating_system", ""),
                endpoint.get("os_version", ""),
                endpoint.get("architecture", ""),
                endpoint.get("processor", ""),
                endpoint.get("ram_gb", ""),
                endpoint.get("disk_total_gb", ""),
                endpoint.get("disk_free_gb", ""),
                endpoint.get("disk_used_percent", ""),
                yes_no(
                    defender.get(
                        "antivirus_enabled"
                    )
                ),
                yes_no(
                    defender.get(
                        "realtime_protection"
                    )
                ),
                format_defender_signature(
                    defender.get(
                        "signature_updated",
                        ""
                    )
                ),
                get_firewall_status(
                    raw_data,
                    "Domain"
                ),
                get_firewall_status(
                    raw_data,
                    "Private"
                ),
                get_firewall_status(
                    raw_data,
                    "Public"
                ),
                updates.get(
                    "pending_updates",
                    ""
                ),
                yes_no(
                    office.get(
                        "installed"
                    )
                ),
                office.get(
                    "product",
                    ""
                ),
                office.get(
                    "version",
                    ""
                ),
                office.get(
                    "platform",
                    ""
                ),
                get_enabled_users(
                    raw_data
                ),
                get_local_administrators(
                    raw_data
                ),
                summary.get(
                    "pass",
                    0
                ),
                summary.get(
                    "review",
                    0
                ),
                summary.get(
                    "fail",
                    0
                ),
                summary.get(
                    "total",
                    0
                ),
                get_overall_status(
                    summary
                ),
                agent.get(
                    "collected_at_utc",
                    ""
                )
            ]
        )

    # ========================================================
    # ASSESSMENT SUMMARY
    # ========================================================

    summary_sheet.append(
        [
            "Hostname",
            "PASS",
            "REVIEW",
            "FAIL",
            "Total Checks",
            "Overall Status"
        ]
    )

    for report in reports:
        endpoint = report.get("endpoint", {})
        summary = report.get("summary", {})

        summary_sheet.append(
            [
                endpoint.get(
                    "hostname",
                    ""
                ),
                summary.get(
                    "pass",
                    0
                ),
                summary.get(
                    "review",
                    0
                ),
                summary.get(
                    "fail",
                    0
                ),
                summary.get(
                    "total",
                    0
                ),
                get_overall_status(
                    summary
                )
            ]
        )

    # ========================================================
    # FINDINGS
    # ========================================================

    findings_sheet.append(
        [
            "Hostname",
            "Status",
            "Category",
            "Check",
            "Current Value",
            "Expected",
            "Recommendation"
        ]
    )

    for report in reports:
        hostname = report.get(
            "endpoint",
            {}
        ).get(
            "hostname",
            ""
        )

        for finding in report.get(
            "findings",
            []
        ):
            findings_sheet.append(
                [
                    hostname,
                    finding.get(
                        "status",
                        ""
                    ),
                    finding.get(
                        "category",
                        ""
                    ),
                    finding.get(
                        "check",
                        ""
                    ),
                    finding.get(
                        "current_value",
                        ""
                    ),
                    finding.get(
                        "expected",
                        ""
                    ),
                    finding.get(
                        "recommendation",
                        ""
                    )
                ]
            )

    # ========================================================
    # LOCAL USERS
    # ========================================================

    users_sheet.append(
        [
            "Hostname",
            "User Name",
            "Enabled",
            "Password Required",
            "Administrator"
        ]
    )

    for report in reports:
        hostname = report.get(
            "endpoint",
            {}
        ).get(
            "hostname",
            ""
        )

        raw_data = report.get(
            "raw_data",
            {}
        )

        local_user_data = raw_data.get(
            "local_users",
            {}
        )

        administrators = local_user_data.get(
            "administrators",
            []
        )

        administrator_names = {
            admin.get(
                "Name",
                ""
            ).split("\\")[-1].lower()
            for admin in administrators
            if admin.get("Name")
        }

        for user in local_user_data.get(
            "users",
            []
        ):
            username = user.get(
                "Name",
                ""
            )

            users_sheet.append(
                [
                    hostname,
                    username,
                    yes_no(
                        user.get(
                            "Enabled"
                        )
                    ),
                    yes_no(
                        user.get(
                            "PasswordRequired"
                        )
                    ),
                    (
                        "Yes"
                        if username.lower()
                        in administrator_names
                        else "No"
                    )
                ]
            )

    # ========================================================
    # GENERAL FORMATTING
    # ========================================================

    for sheet in [
        inventory_sheet,
        summary_sheet,
        findings_sheet,
        users_sheet
    ]:
        format_sheet(sheet)

    # ========================================================
    # INVENTORY COLORING
    # ========================================================

    for row in inventory_sheet.iter_rows(
        min_row=2
    ):
        # Operating System
        apply_os_color(row[1])

        # RAM
        apply_ram_color(row[5])

        # Disk Used %
        apply_disk_color(row[8])

        # Defender + Real-Time
        apply_yes_no_color(row[9])
        apply_yes_no_color(row[10])

        # Firewall Domain/Private/Public
        apply_enabled_disabled_color(row[12])
        apply_enabled_disabled_color(row[13])
        apply_enabled_disabled_color(row[14])

        # Pending Updates
        apply_updates_color(row[15])

        # Office Installed
        apply_yes_no_color(
            row[16],
            no_is_review=True
        )

        # PASS / REVIEW / FAIL counts
        if isinstance(row[22].value, int):
            if row[22].value > 0:
                fill_cell(
                    row[22],
                    PASS_FILL
                )

        if isinstance(row[23].value, int):
            if row[23].value > 0:
                fill_cell(
                    row[23],
                    REVIEW_FILL
                )

        if isinstance(row[24].value, int):
            if row[24].value > 0:
                fill_cell(
                    row[24],
                    FAIL_FILL
                )

        # Overall Status
        apply_status_color(row[26])

    # ========================================================
    # SUMMARY COLORING
    # ========================================================

    for row in summary_sheet.iter_rows(
        min_row=2
    ):
        if isinstance(row[1].value, int):
            if row[1].value > 0:
                fill_cell(
                    row[1],
                    PASS_FILL
                )

        if isinstance(row[2].value, int):
            if row[2].value > 0:
                fill_cell(
                    row[2],
                    REVIEW_FILL
                )

        if isinstance(row[3].value, int):
            if row[3].value > 0:
                fill_cell(
                    row[3],
                    FAIL_FILL
                )

        apply_status_color(
            row[5]
        )

    # ========================================================
    # FINDINGS COLORING
    # ========================================================

    for row in findings_sheet.iter_rows(
        min_row=2
    ):
        apply_status_color(
            row[1]
        )

    # ========================================================
    # LOCAL USER COLORING
    # ========================================================

    for row in users_sheet.iter_rows(
        min_row=2
    ):
        enabled = str(
            row[2].value
        ).strip().lower()

        password_required = str(
            row[3].value
        ).strip().lower()

        administrator = str(
            row[4].value
        ).strip().lower()

        if enabled == "yes":
            fill_cell(
                row[2],
                PASS_FILL
            )

        else:
            fill_cell(
                row[2],
                NEUTRAL_FILL
            )

        if (
            enabled == "yes"
            and password_required == "no"
        ):
            fill_cell(
                row[3],
                REVIEW_FILL
            )

        elif password_required == "yes":
            fill_cell(
                row[3],
                PASS_FILL
            )

        else:
            fill_cell(
                row[3],
                NEUTRAL_FILL
            )

        if (
            enabled == "yes"
            and administrator == "yes"
        ):
            fill_cell(
                row[4],
                REVIEW_FILL
            )

        elif administrator == "yes":
            fill_cell(
                row[4],
                NEUTRAL_FILL
            )

    # ========================================================
    # SAVE
    # ========================================================

    workbook.save(
        OUTPUT_FILE
    )


if __name__ == "__main__":
    reports = load_reports()

    if not reports:
        print(
            "No endpoint assessment files "
            "found inside endpoint-data."
        )

    else:
        build_excel(
            reports
        )

        print()
        print("=" * 60)
        print("NANOBRICKS MASTER ENDPOINT REPORT")
        print("=" * 60)

        print(
            f"Endpoints processed: "
            f"{len(reports)}"
        )

        print(
            f"Excel report: "
            f"{OUTPUT_FILE}"
        )

        print("=" * 60)
